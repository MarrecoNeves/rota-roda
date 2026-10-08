"""Gera app/static/modelo_pontos.xlsx (planilha-modelo do Rota Roda).

Uso:  pip install openpyxl  &&  python tools/gerar_modelo.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "app" / "static" / "modelo_pontos.xlsx"

AZUL, AZUL_ESC, CINZA, AZUL_CLARO = "004F9F", "002F5F", "EAEAEA", "DCE9F7"
BRANCO = "FFFFFF"
FIRST, LAST = 5, 204  # 200 linhas prontas para colar

COLS = [  # (cabeçalho, largura, obrigatória)
    ("Nome (opcional)", 28, False),
    ("Endereço (rua, nº, bairro, cidade)", 60, True),
    ("CEP (opcional)", 15, False),
    ("Demanda", 12, False),
    ("Depósito?", 12, False),
    ("Abre às", 11, False),
    ("Fecha às", 11, False),
    ("Atendimento (min)", 15, False),
]

EXEMPLO = [
    ("Escola de Engenharia UFF", "Rua Passo da Pátria, 156, São Domingos, Niterói - RJ", "", 0, "Sim", None, None, None),
    ("Presidente Backer", "Rua Presidente Backer, 337, Icaraí, Niterói - RJ", "", 100, "Não", None, None, 5),
    ("Conceição", "Rua da Conceição, 100, Centro, Niterói - RJ", "", 80, "Não", None, None, 5),
    ("Roberto Silveira", "Avenida Roberto Silveira, 512, Icaraí, Niterói - RJ", "", 70, "Não", None, None, 5),
    ("Nóbrega", "Rua Nóbrega, 672, Icaraí, Niterói - RJ", "", 80, "Não", None, None, 5),
    ("Marechal Deodoro", "Rua Marechal Deodoro, 200, Centro, Niterói - RJ", "", 100, "Não", None, None, 5),
    ("Leite Ribeiro", "Rua Leite Ribeiro, 212, Fonseca, Niterói - RJ", "", 50, "Não", None, None, 5),
    ("Tenente Osório", "Rua Tenente Osório, 30, Fonseca, Niterói - RJ", "", 60, "Não", None, None, 5),
]

thin = Side(style="thin", color="C9D3E0")
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)


def faixa(ws, row, text, *, size=11, bold=False, color="1F2937", fill=None, height=None, italic=False):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLS))
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=color)
    c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    if fill:
        for col in range(1, len(COLS) + 1):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height


def folha_pontos(ws, titulo, subtitulo, linhas, nome_tabela):
    ws.sheet_view.showGridLines = False
    faixa(ws, 1, titulo, size=16, bold=True, color=BRANCO, fill=AZUL, height=36)
    faixa(ws, 2, subtitulo, size=10.5, fill=CINZA, height=34)
    ws.row_dimensions[3].height = 8

    for j, (head, width, req) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = width
        c = ws.cell(row=4, column=j, value=head)
        c.font = Font(bold=True, color=BRANCO, size=11)
        c.fill = PatternFill("solid", fgColor=AZUL_ESC if req else AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[4].height = 34

    for i in range(FIRST, LAST + 1):
        ws.row_dimensions[i].height = 20
        vals = linhas[i - FIRST] if i - FIRST < len(linhas) else None
        for j in range(1, len(COLS) + 1):
            c = ws.cell(row=i, column=j)
            if vals is not None and vals[j - 1] not in (None, ""):
                c.value = vals[j - 1]
            c.border = BORDA
            c.alignment = Alignment(vertical="center", horizontal="left" if j <= 3 else "center")
            if j == 3:
                c.number_format = "@"          # CEP como texto: não perde o zero da frente
            elif j in (6, 7):
                c.number_format = "hh:mm"
            elif j in (4, 8):
                c.number_format = "0"

    ref = f"A4:{get_column_letter(len(COLS))}{LAST}"
    tab = Table(displayName=nome_tabela, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)

    # Linha do depósito destacada
    ws.conditional_formatting.add(
        f"A{FIRST}:{get_column_letter(len(COLS))}{LAST}",
        FormulaRule(formula=[f'LOWER($E{FIRST})="sim"'], fill=PatternFill("solid", fgColor=AZUL_CLARO),
                    font=Font(bold=True, color=AZUL_ESC)),
    )

    def dv(rng, **kw):
        v = DataValidation(allow_blank=True, showInputMessage=True, showErrorMessage=True, **kw)
        ws.add_data_validation(v)
        v.add(rng)

    dv(f"B{FIRST}:B{LAST}", type="textLength", operator="greaterThanOrEqual", formula1="0",
       promptTitle="Endereço", prompt="Rua, número, bairro e cidade.\nEx.: Rua da Conceição, 100, Centro, Niterói - RJ")
    dv(f"C{FIRST}:C{LAST}", type="textLength", operator="lessThanOrEqual", formula1="9",
       promptTitle="CEP (opcional)", prompt="Use se não tiver o endereço completo.\nEx.: 24020-085",
       errorTitle="CEP", error="Digite o CEP com 8 números, ex.: 24020-085")
    dv(f"D{FIRST}:D{LAST}", type="whole", operator="greaterThanOrEqual", formula1="0",
       promptTitle="Demanda", prompt="Quanto entregar neste ponto (número inteiro).\nDeixe 0 no depósito.",
       errorTitle="Demanda", error="Use um número inteiro, ex.: 80")
    dv(f"E{FIRST}:E{LAST}", type="list", formula1='"Sim,Não"',
       promptTitle="Depósito?", prompt="Sim = ponto de saída dos veículos.\nSe nenhum for marcado, a 1ª linha vira o depósito.",
       errorTitle="Depósito?", error="Escolha Sim ou Não")
    for col, nome in (("F", "Abre às"), ("G", "Fecha às")):
        dv(f"{col}{FIRST}:{col}{LAST}", type="time", operator="between", formula1="0", formula2="0.999988",
           promptTitle=nome, prompt="Opcional. Horário no formato 09:00.\nPreencher ativa as janelas de tempo.",
           errorTitle=nome, error="Use o formato de hora, ex.: 09:00")
    dv(f"H{FIRST}:H{LAST}", type="whole", operator="between", formula1="0", formula2="600",
       promptTitle="Atendimento", prompt="Opcional. Minutos parado no cliente (padrão 5).",
       errorTitle="Atendimento", error="Use minutos inteiros, ex.: 10")

    ws.freeze_panes = f"A{FIRST}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"


def folha_como_usar(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 52
    ws.column_dimensions["E"].width = 40
    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = "Como usar a planilha do Rota Roda"
    t.font = Font(size=16, bold=True, color=BRANCO)
    t.alignment = Alignment(vertical="center", indent=1)
    for col in "ABCDE":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=AZUL)
    ws.row_dimensions[1].height = 36

    passos = [
        "Abra a aba Pontos. Na linha 5 já está o depósito (ponto de saída): cole o endereço dele na coluna Endereço.",
        "Da linha 6 para baixo, cole os endereços dos clientes, um por linha. Pode copiar uma coluna inteira de outra planilha.",
        "Preencha a Demanda de cada cliente. O resto é opcional.",
        "Salve e, no app, clique em Importar planilha (ou arraste o arquivo para a tela).",
    ]
    r = 3
    ws.cell(row=r, column=2, value="Passo a passo").font = Font(size=12, bold=True, color=AZUL)
    for i, p in enumerate(passos, start=1):
        r += 1
        n = ws.cell(row=r, column=1, value=i)
        n.font = Font(bold=True, color=BRANCO)
        n.fill = PatternFill("solid", fgColor=AZUL)
        n.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2, value=p)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[r].height = 30

    r += 2
    ws.cell(row=r, column=2, value="Colunas").font = Font(size=12, bold=True, color=AZUL)
    r += 1
    for j, h in enumerate(["Coluna", "Obrigatória?", "O que colocar", "Exemplo"], start=2):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(bold=True, color=BRANCO)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.border = BORDA
        c.alignment = Alignment(vertical="center", indent=1)
    linhas = [
        ("Endereço", "Sim", "Rua, número, bairro e cidade. Quanto mais completo, melhor.", "Rua da Conceição, 100, Centro, Niterói - RJ"),
        ("Demanda", "Recomendada", "Quanto entregar no ponto (número inteiro). Vazio = 10.", "80"),
        ("Depósito?", "Não", "Sim no ponto de saída. Sem nenhum Sim, a 1ª linha vira o depósito.", "Sim"),
        ("Nome", "Não", "Como o ponto aparece no mapa e no relatório.", "Loja Centro"),
        ("CEP", "Não", "Só quando não tiver o endereço completo.", "24020-085"),
        ("Abre às / Fecha às", "Não", "Janela de atendimento. Preencher ativa as janelas de tempo.", "09:00 / 12:00"),
        ("Atendimento (min)", "Não", "Minutos parado no cliente.", "10"),
    ]
    for k, row in enumerate(linhas):
        r += 1
        for j, v in enumerate(row, start=2):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDA
            c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            if k % 2:
                c.fill = PatternFill("solid", fgColor="F4F7FB")
        ws.cell(row=r, column=2).font = Font(bold=True)
        ws.row_dimensions[r].height = 30

    r += 2
    ws.cell(row=r, column=2, value="Dicas").font = Font(size=12, bold=True, color=AZUL)
    dicas = [
        "Não quer usar planilha? No app, o botão Colar endereços aceita uma lista copiada de qualquer lugar.",
        "Sua própria planilha também serve: o app reconhece colunas como Endereço, Rua + Número + Bairro + Cidade, CEP ou Latitude/Longitude, mesmo com título em cima.",
        "A aba Exemplo mostra um caso pronto em Niterói. Se a aba Pontos estiver vazia, o app importa o Exemplo.",
        "O mapa gratuito localiza cerca de 1 endereço por segundo: 30 endereços levam por volta de meio minuto.",
    ]
    for d in dicas:
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2, value="•  " + d)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[r].height = 30
    r += 2
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = ws.cell(row=r, column=2, value="Rota Roda · Daniel Neves, Juan Souza, Pedro Souza e Pedro Jensen · UFF")
    c.font = Font(italic=True, color="6B7280", size=9)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pontos"
    folha_pontos(
        ws, "Rota Roda · Pontos de entrega",
        "Cole os endereços na coluna Endereço, um por linha. A linha 5 é o depósito (ponto de saída). "
        "Só o endereço é obrigatório; a Demanda é recomendada.",
        [("Depósito (ponto de saída)", "", "", 0, "Sim", None, None, None)],
        "Pontos",
    )
    ws.sheet_properties.tabColor = AZUL
    ex = wb.create_sheet("Exemplo")
    folha_pontos(
        ex, "Exemplo pronto · Niterói",
        "Só para consulta: o app usa a aba Pontos. Se a aba Pontos estiver vazia, ele importa este exemplo.",
        EXEMPLO, "Exemplo",
    )
    ex.sheet_properties.tabColor = "7FA7D1"
    cu = wb.create_sheet("Como usar")
    folha_como_usar(cu)
    cu.sheet_properties.tabColor = "9CA3AF"
    wb.active = 0
    ws.sheet_view.selection[0].activeCell = "B5"
    ws.sheet_view.selection[0].sqref = "B5"
    wb.properties.creator = "Rota Roda (UFF)"
    wb.properties.title = "Rota Roda - modelo de pontos"
    wb.save(OUT)
    print("ok", OUT)


if __name__ == "__main__":
    main()
